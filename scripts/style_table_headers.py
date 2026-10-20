#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Post-process Word tables: center and bold the first row of every table.
"""

import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL


def style_table_headers(docx_path):
    doc = Document(docx_path)
    for table in doc.tables:
        if not table.rows:
            continue
        header_row = table.rows[0]
        for cell in header_row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
    doc.save(docx_path)


def main():
    if len(sys.argv) != 2:
        print("Usage: python style_table_headers.py <docx_path>", file=sys.stderr)
        return 2
    style_table_headers(sys.argv[1])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
