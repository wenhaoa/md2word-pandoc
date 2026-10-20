"""Isolated regressions for content preservation and the complete converter."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import redirect_stdout
from io import StringIO
import base64
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

import add_captions
import autofit_tables
import equation_numbering
import fix_heading_numbering
import inline_math_to_text
import merge_cover
import merge_math_runs
import validate_docx
import xref_fields
import xref_sections
from docx_safety import bookmark

HERE = Path(__file__).resolve().parent
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def math_element(inner, display=False):
    xml = f'<m:oMath xmlns:m="{M}">{inner}</m:oMath>'
    if display:
        xml = f'<m:oMathPara xmlns:m="{M}">{xml}</m:oMathPara>'
    return parse_xml(xml)


class WordTests(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory(prefix='md2word-regression-')
        self.root = Path(self.scratch.name)

    def tearDown(self):
        self.scratch.cleanup()

    def save(self, doc, name='fixture.docx'):
        path = self.root / name
        doc.save(path)
        return path

    def run_node(self, name, args, env=None):
        return subprocess.run(['node', str(HERE / name), *map(str, args)], capture_output=True,
                              text=True, encoding='utf-8', env={**os.environ, 'PYTHONIOENCODING': 'utf-8', **(env or {})}, timeout=90)

    def fixture(self, name='temp_input.md'):
        # One real bitmap, referenced twice with deliberately identical captions.
        image = self.root / '试验 图.png'
        image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2nV8AAAAASUVORK5CYII='))
        source = self.root / name
        source.write_text('# 测试 & "标题"\n\n## 1 概述\n\n### 1.1 输入\n\n**表1 参数**\n\n| 参数 | 值 |\n| --- | --- |\n| 推力 | 20.1234mN |\n| 图1和图2 | 见表1 |\n\n![试验图](<试验 图.png>)\n\n**图1 同名工况图**\n\n![另一幅图](<试验 图.png>)\n\n**图2 同名工况图**\n\n见图1、图2、表1、第1章、第1.1节。\n\n## 2 计算\n\n2026年任务，3D模型，300km离轨标准。\n\n$$\n\\frac{a}{b}=c\\tag{2-1}\n$$\n\n按式(2-1)计算。行内$x_{\\frac{a}{b}}$保留。\n\n`工具 --名称 中文`；`数组[1]`。\n\n~~~\n值 = 5 mm\n\\tag{99-99}\n~~~\n', encoding='utf-8')
        return source

    def test_cover_never_removes_drawing_math_field_or_break(self):
        for tag in ('w:drawing', 'w:fldChar', 'w:br', 'w:bookmarkStart', 'math'):
            doc = Document()
            doc.add_paragraph()._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
            p = doc.add_paragraph()
            if tag == 'math':
                p._p.append(math_element('<m:r><m:t>x</m:t></m:r>', True))
            else:
                p.add_run()._r.append(OxmlElement(tag))
            before = doc.element.body.xml
            merge_cover._remove_first_empty_paragraph(doc)
            self.assertEqual(doc.element.body.xml, before, tag)
        doc = Document()
        doc.add_paragraph()._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        doc.add_paragraph()
        doc.add_paragraph('正文')
        merge_cover._remove_first_empty_paragraph(doc)
        self.assertEqual(len(doc.paragraphs), 2)

    def test_nested_fraction_and_nested_script_are_preserved(self):
        for sub in ('<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>',
                    '<m:sSup><m:e><m:r><m:t>a</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'):
            doc = Document()
            om = math_element('<m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub>' + sub + '</m:sub></m:sSub>')
            doc.add_paragraph()._p.append(om)
            before = doc.element.body.xml
            self.assertEqual(inline_math_to_text.convert(doc), (0, 1))
            self.assertEqual(doc.element.body.xml, before)

    def test_composition_preserves_nested_bookmark_pairs(self):
        prefix = Document()
        cover = prefix.add_paragraph('封面')
        bookmark(cover._p, cover.runs[0]._r, cover.runs[0]._r, 'Cover', 0)
        prefix_path = self.save(prefix, 'cover.docx')
        body = Document()
        first = body.add_heading('概述', level=1)
        caption = body.add_paragraph('图1-1 测试题注')
        # Pandoc creates an outer section bookmark around body content; the
        # generated caption bookmark is nested inside it. Cover IDs may collide.
        bookmark(body.element.body, first._p, caption._p, 'pandoc_section', 0)
        add_captions.process_captions(body)
        body_path = self.save(body, 'body.docx')
        output = self.root / 'composed.docx'
        merge_cover.merge(prefix_path, body_path, output)
        validate_docx.validate(output)
        merged = Document(output)
        starts = {element.get(qn('w:name')): element
                  for element in merged.element.body.iter(qn('w:bookmarkStart'))}
        ends = {element.get(qn('w:id')): element
                for element in merged.element.body.iter(qn('w:bookmarkEnd'))}
        start = starts['Fig_1_1']
        end = ends[start.get(qn('w:id'))]
        self.assertIs(start.getparent(), end.getparent())
        self.assertEqual(''.join(element.text or '' for element in start.getparent().iter(qn('w:t'))),
                         '图1-1 测试题注')

    def test_simple_inline_math_can_convert(self):
        doc = Document()
        doc.add_paragraph()._p.append(math_element('<m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>'))
        self.assertEqual(inline_math_to_text.convert(doc), (1, 0))
        self.assertEqual(doc.paragraphs[0].text, 'xi')
        self.assertTrue(doc.paragraphs[0].runs[1].font.subscript)

    def test_numeric_heading_and_caption_title_are_preserved(self):
        doc = Document()
        for title in ('2026年任务', '3D模型', '300km离轨标准'):
            doc.add_heading(title, level=1)
        doc.add_paragraph('图1 300km离轨标准')
        path = self.save(doc)
        fix_heading_numbering.run(path)
        self.assertEqual([p.text for p in Document(path).paragraphs[:3]], ['2026年任务', '3D模型', '300km离轨标准'])
        doc = Document(path)
        add_captions.process_captions(doc)
        self.assertIn('300km离轨标准', doc.paragraphs[-1].text)

    def test_keep_manual_flag(self):
        doc = Document()
        doc.add_heading('1 概述', level=1)
        path = self.save(doc)
        fix_heading_numbering.run(path, keep_manual=True)
        self.assertEqual(Document(path).paragraphs[0].text, '1 概述')

    def test_duplicate_titles_belonging_to_different_images_are_retained(self):
        doc = Document()
        doc.add_heading('概述', level=1)
        for number in (1, 2):
            doc.add_paragraph().add_run()._r.append(OxmlElement('w:drawing'))
            doc.add_paragraph(f'图1-{number} 同名工况图')
        self.assertEqual(add_captions.process_captions(doc, dedup=True), 2)
        self.assertEqual(len([p for p in doc.paragraphs if '同名工况图' in p.text]), 2)
        self.assertEqual(add_captions.process_captions(doc), 0)

    def test_unnumbered_caption_does_not_jump_over_body_text(self):
        doc = Document()
        doc.add_paragraph().add_run()._r.append(OxmlElement('w:drawing'))
        doc.add_paragraph('这是正文。')
        doc.add_paragraph('图 系统分析')
        self.assertEqual(add_captions.process_captions(doc), 0)

    def test_body_table_before_first_chapter_is_formatted(self):
        doc = Document()
        doc.add_table(rows=2, cols=2)
        doc.add_heading('概述', level=1)
        path = self.save(doc)
        output = StringIO()
        with redirect_stdout(output):
            autofit_tables.run(path, body_only=True)
        self.assertIn('处理表格 1 张', output.getvalue())

    def test_equation_plan_is_read_and_rerun_is_idempotent(self):
        doc = Document()
        for number in range(1, 5):
            doc.add_heading('章节', level=1)
        doc.add_paragraph()._p.append(math_element('<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>', True))
        doc.add_paragraph('见式(4-1)。')
        path = self.save(doc)
        cli = subprocess.run([sys.executable, '-X', 'utf8', str(HERE / 'equation_numbering.py'), str(path), '--plan', '4-1'], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(cli.returncode, 0, cli.stderr)
        equation_numbering.run(path, plan=['4-1'])
        doc = Document(path)
        instructions = [e.text or '' for e in doc.element.body.iter(qn('w:instrText'))]
        self.assertEqual(sum('SEQ 公式' in value for value in instructions), 1)
        self.assertEqual(sum('REF Eq_4_1' in value for value in instructions), 1)
        equation = next(p for p in doc.paragraphs if equation_numbering.has_display_math(p._p))
        self.assertEqual(equation._p.find(qn('w:pPr')).find(qn('w:jc')).get(qn('w:val')), 'left')
        validate_docx.validate(path)

    def test_bad_plan_aborts_without_saving(self):
        doc = Document()
        doc.add_heading('计算', level=1)
        doc.add_paragraph()._p.append(math_element('<m:r><m:t>x</m:t></m:r>', True))
        path = self.save(doc)
        before = path.read_bytes()
        with self.assertRaises(ValueError):
            equation_numbering.run(path, plan=['4-7'])
        self.assertEqual(path.read_bytes(), before)
        with self.assertRaises(ValueError):
            equation_numbering.run(path, plan=[])

    def test_caption_parsing_accepts_plain_bold_and_global(self):
        path = self.root / 'captions.md'
        path.write_text('图1-1 标题\n**图1-2 标题**\n表1\u3000参数\n~~~\n图8-1 示例\n~~~\n', encoding='utf-8')
        self.assertEqual(xref_fields.md_caption_numbers(path), (['1-1', '1-2'], ['1']))

    def test_math_merging_preserves_word_properties_and_code(self):
        doc = Document()
        p = doc.add_paragraph('代码5 mm，文件路径。')
        om = math_element('<m:r><m:t>x</m:t></m:r><m:r><m:t>y</m:t></m:r>')
        props = OxmlElement('w:rPr')
        props.append(OxmlElement('w:b'))
        om[1].insert(0, props)
        p._p.append(om)
        self.assertEqual(merge_math_runs.merge_in(om), 0)
        merge_math_runs.normalize_spaces(doc)
        self.assertIn('5 mm', p.text)

    def test_full_conversion_preserves_source_sentinels_and_references(self):
        source = self.fixture()
        sentinel = self.root / 'temp_output.docx'
        sentinel.write_text('KEEP ORIGINAL', encoding='utf-8')
        original = source.read_bytes()
        result = self.run_node('run_conversion.js', [source])
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(sentinel.read_text(encoding='utf-8'), 'KEEP ORIGINAL')
        outputs = list(self.root.glob('temp_input[0-9]*.docx'))
        self.assertEqual(len(outputs), 1)
        self.assertEqual(list(self.root.glob('.md2word-*')), [])
        doc = Document(outputs[0])
        text = '\n'.join(p.text for p in doc.paragraphs)
        for value in ('同名工况图', '工具 --名称 中文', '值 = 5 mm', '\\tag{99-99}', '2026年任务，3D模型，300km离轨标准。'):
            self.assertIn(value, text)
        instructions = ''.join(e.text or '' for e in doc.element.body.iter(qn('w:instrText')))
        for name in ('Fig_1', 'Fig_2', 'Tbl_1', 'Eq_2_1', 'Sec_1', 'Sec_1_1'):
            self.assertIn('REF ' + name, instructions)
        self.assertFalse(any(element.get(qn('w:dirty')) == 'true'
                             for element in doc.element.body.iter(qn('w:fldChar'))),
                         'Valid cached fields must not trigger an update prompt on open')
        self.assertEqual(len(list(doc.element.body.iter('{' + M + '}f'))), 2)
        cover = Document(HERE.parent / 'templates' / 'md2word模板.docx')
        cover_drawings = len(list(cover.element.body.iter(qn('w:drawing'))))
        self.assertEqual(len(list(doc.element.body.iter(qn('w:drawing')))), 2 + cover_drawings)
        self.assertEqual(sum('同名工况图' in p.text for p in doc.paragraphs), 2)

    def test_concurrent_conversions_have_unique_outputs(self):
        source = self.fixture('concurrent.md')
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: self.run_node('run_conversion.js', [source]), range(2)))
        self.assertTrue(all(result.returncode == 0 for result in results), '\n'.join(result.stdout + result.stderr for result in results))
        self.assertEqual(len(list(self.root.glob('concurrent[0-9]*.docx'))), 2)

    def test_critical_stage_failure_is_not_reported_as_success(self):
        source = self.fixture('failed.md')
        isolated = self.root / 'isolated-skill'
        shutil.copytree(HERE, isolated / 'scripts', ignore=shutil.ignore_patterns('__pycache__'))
        shutil.copytree(HERE.parent / 'templates', isolated / 'templates')
        (isolated / 'scripts' / 'equation_numbering.py').write_text('raise RuntimeError("injected failure")\n', encoding='utf-8')
        result = self.run_node('run_conversion.js', [source], {'MD2WORD_SKILL_DIR': str(isolated)})
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn('转换及结构检查通过。输出文件', result.stdout)
        self.assertEqual(list(self.root.glob('failed[0-9]*.docx')), [])
        self.assertTrue(list(self.root.glob('.md2word-*')))

    def test_retired_migrations_cannot_rewrite_installed_files(self):
        protected = [HERE.parent / 'SKILL.md', HERE / 'run_conversion.js', HERE / 'add_captions.py']
        before = [hashlib.sha256(path.read_bytes()).hexdigest() for path in protected]
        for name in ('update_skill_doc', 'patch_caption_pipeline'):
            result = subprocess.run([sys.executable, '-X', 'utf8', str(HERE / (name + '.py'))], capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 2)
            result = subprocess.run([sys.executable, '-c', 'import ' + name], cwd=HERE, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual([hashlib.sha256(path.read_bytes()).hexdigest() for path in protected], before)

    def test_preflight_accepts_existing_headings_bold_captions_and_refs(self):
        result = self.run_node('check_markdown.js', [self.fixture(), '--json'])
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(json.loads(result.stdout)['issues'], [])

    def test_legacy_alt_caption_and_cleanup_remain_available(self):
        self.fixture()
        source = self.root / 'legacy.md'
        body = '中文 A，温度 5 mm，"中文"。'
        content = ('---\ntitle: 兼容性检查\n---\n\n## 1. 概述\n\n' + body +
                   '\n\n![图1.1 工况图](<试验 图.png>)\n\n见图1.1。\n\n`代码 5 mm "中文"`\n')
        source.write_text(content, encoding='utf-8')
        check = self.run_node('check_markdown.js', [source, '--json'])
        self.assertEqual(check.returncode, 0, check.stdout + check.stderr)
        result = self.run_node('run_conversion.js', [source])
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        output = next(self.root.glob('legacy[0-9]*.docx'))
        doc = Document(output)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('中文A，温度5mm，“中文”。', text)
        self.assertIn('代码 5 mm "中文"', text)
        self.assertEqual(sum('工况图' in p.text for p in doc.paragraphs), 1)
        instructions = ''.join(e.text or '' for e in doc.element.body.iter(qn('w:instrText')))
        self.assertIn('REF Fig_1_1', instructions)
        self.assertEqual(source.read_text(encoding='utf-8'), content)
        previous = set(self.root.glob('legacy[0-9]*.docx'))
        result = self.run_node('run_conversion.js', [source, '--preserve-text'])
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        preserved = (set(self.root.glob('legacy[0-9]*.docx')) - previous).pop()
        self.assertIn(body, '\n'.join(p.text for p in Document(preserved).paragraphs))

    def test_preflight_reports_mermaid_but_ignores_ordinary_code(self):
        source = self.root / 'diagram.md'
        source.write_text('```mermaid\ngraph TD\nA-->B\n```\n\n```text\n图9-9 示例\n```\n', encoding='utf-8')
        result = self.run_node('check_markdown.js', [source, '--json'])
        self.assertEqual(result.returncode, 1)
        self.assertEqual([issue['kind'] for issue in json.loads(result.stdout)['issues']], ['Mermaid'])

    def test_paths_with_shell_characters_and_chapter_title_fallback(self):
        source = self.root / "工程! & '测试'.md"
        source.write_text('# 1 概述\n\n正文。\n\n# 2 计算\n\n结果。\n', encoding='utf-8')
        result = self.run_node('run_conversion.js', [source])
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        output = next(self.root.glob(source.stem + '[0-9]*.docx'))
        text = '\n'.join(p.text for p in Document(output).paragraphs)
        self.assertIn(source.stem, text)
        self.assertTrue(source.exists())

    def test_lua_preserves_real_titles_and_handles_chinese_ordinals(self):
        for heading in ('二、概述', '3.14km测量', '五．概述', '三维、模型'):
            source = self.root / 'heading.md'
            source.write_text('# ' + heading + '\n', encoding='utf-8')
            result = subprocess.run(['pandoc', str(source), '--to', 'json', '-M', 'md2word_heading_mode=chapters', '--lua-filter', str(HERE / 'style_filter.lua')], capture_output=True, text=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            # The Lua filter logs before its JSON output.
            ast = json.loads(result.stdout[result.stdout.index('{'):])
            if heading in ('二、概述', '五．概述'):
                self.assertEqual(ast['blocks'][0]['c'][2][0]['c'], '概述')
            else:
                self.assertEqual(ast['blocks'][0]['c'][2][0]['c'], heading)


if __name__ == '__main__':
    unittest.main(verbosity=2)
