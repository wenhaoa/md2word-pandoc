"""Regressions for explicit landscape layout, restoration, and source safety."""
from contextlib import redirect_stdout
from io import StringIO
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import landscape_tables
from table_sections import section_for_block

HERE = Path(__file__).resolve().parent


def orientation(section):
    return section.find(qn('w:pgSz')).get(qn('w:orient'), 'portrait')


class LandscapeTests(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory(prefix='md2word-landscape-')
        self.root = Path(self.scratch.name)

    def tearDown(self):
        self.scratch.cleanup()

    def table(self, doc, text):
        table = doc.add_table(rows=2, cols=2)
        for i, row in enumerate(table.rows):
            row.cells[0].text = text + str(i)
            row.cells[1].text = 't0+1259.76'
        return table

    def apply(self, doc, selected, count):
        path = self.root / 'test.docx'
        doc.save(path)
        with redirect_stdout(StringIO()):
            landscape_tables.apply(path, {'table_count': count, 'landscape': selected})
        return Document(path)

    def test_only_selected_table_rotates_and_following_table_stays_portrait(self):
        doc = Document()
        doc.add_paragraph('前文')
        self.table(doc, '横向')
        doc.add_paragraph('后文')
        self.table(doc, '竖向')
        before = [[cell.text for row in table.rows for cell in row.cells] for table in doc.tables]
        margins = (doc.sections[0].left_margin, doc.sections[0].right_margin)
        result = self.apply(doc, [0], 2)
        self.assertEqual([orientation(s._sectPr) for s in result.sections], ['portrait', 'landscape', 'portrait'])
        self.assertEqual(before, [[cell.text for row in table.rows for cell in row.cells] for table in result.tables])
        self.assertEqual([orientation(section_for_block(result, t._tbl)) for t in result.tables], ['landscape', 'portrait'])
        self.assertTrue(all((s.left_margin, s.right_margin) == margins for s in result.sections))
        self.assertIsNotNone(result.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')))
        grid = result.tables[0]._tbl.find(qn('w:tblGrid'))
        total = sum(int(col.get(qn('w:w'))) for col in grid)
        sec = section_for_block(result, result.tables[0]._tbl)
        self.assertLessEqual(total, int(sec.find(qn('w:pgSz')).get(qn('w:w'))) - sum(int(sec.find(qn('w:pgMar')).get(qn('w:' + key))) for key in ('left', 'right')))

    def test_first_last_only_and_adjacent_tables_do_not_add_empty_sections(self):
        for before, after, selected in [(False, True, [0]), (True, False, [0]), (False, False, [0]), (False, False, [0, 1])]:
            doc = Document()
            if before:
                doc.add_paragraph('前文')
            for _ in selected:
                self.table(doc, '横向')
            if after:
                doc.add_paragraph('后文')
            result = self.apply(doc, selected, len(selected))
            expected = (['portrait'] if before else []) + ['landscape'] * len(selected) + (['portrait'] if after else [])
            self.assertEqual([orientation(s._sectPr) for s in result.sections], expected)

    def test_cover_offset_and_existing_section_are_preserved(self):
        doc = Document()
        self.table(doc, '封面')
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph('正文')
        self.table(doc, '目标')
        doc.add_paragraph('尾部')
        cover = doc.tables[0]._tbl.xml
        result = self.apply(doc, [0], 1)
        self.assertEqual(result.tables[0]._tbl.xml, cover)
        self.assertEqual(orientation(section_for_block(result, result.tables[0]._tbl)), 'portrait')
        self.assertEqual(orientation(section_for_block(result, result.tables[1]._tbl)), 'landscape')

    def test_invalid_index_does_not_write(self):
        doc = Document()
        self.table(doc, '原文')
        path = self.root / 'invalid.docx'
        doc.save(path)
        before = path.read_bytes()
        with self.assertRaises(ValueError):
            landscape_tables.apply(path, {'table_count': 1, 'landscape': [1]})
        self.assertEqual(path.read_bytes(), before)

    def test_split_sections_continue_original_page_numbering(self):
        doc = Document()
        doc.add_paragraph('前文')
        self.table(doc, '横向')
        doc.add_paragraph('后文')
        numbering = OxmlElement('w:pgNumType')
        numbering.set(qn('w:start'), '7')
        doc.sections[0]._sectPr.append(numbering)
        result = self.apply(doc, [0], 1)
        self.assertEqual([s._sectPr.find(qn('w:pgNumType')).get(qn('w:start')) for s in result.sections], ['7', None, None])

    def test_full_converter_handles_marked_and_unmarked_tables(self):
        source = self.root / '横向测试.md'
        content = ('---\ntitle: 横向测试\n---\n\n## 1. 正文\n\n前文。\n\n'
                   '::: {.landscape}\n\n| 参数 | 值 |\n| ------ | ---------------- |\n| 时间 | t0+1259.76 |\n\n:::\n\n'
                   '后文。\n\n| 参数 | 值 |\n| ------ | -------- |\n| 类型 | float |\n')
        source.write_text(content, encoding='utf-8')
        result = subprocess.run(['node', str(HERE / 'run_conversion.js'), str(source)], capture_output=True,
                                text=True, encoding='utf-8', env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, timeout=120)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        doc = Document(next(self.root.glob('横向测试[0-9]*.docx')))
        tables = doc.tables[-2:]
        self.assertEqual([orientation(section_for_block(doc, t._tbl)) for t in tables], ['landscape', 'portrait'])
        self.assertEqual(orientation(doc.sections[0]._sectPr), 'portrait')
        self.assertEqual(orientation(doc.sections[-1]._sectPr), 'portrait')
        self.assertEqual(source.read_text(encoding='utf-8'), content)
        self.assertNotIn('landscape', '\n'.join(p.text for p in doc.paragraphs))
        self.assertFalse(list(self.root.glob('.md2word-*')))

    def test_invalid_marker_fails_before_output(self):
        source = self.root / 'invalid.md'
        source.write_text('::: {.landscape}\n\n不是表格。\n\n:::\n', encoding='utf-8')
        result = subprocess.run(['node', str(HERE / 'run_conversion.js'), str(source)], capture_output=True,
                                text=True, encoding='utf-8', env={**os.environ, 'PYTHONIOENCODING': 'utf-8'}, timeout=90)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('landscape块只能包含一张表格', result.stderr)
        self.assertFalse(list(self.root.glob('invalid[0-9]*.docx')))


if __name__ == '__main__':
    unittest.main(verbosity=2)
